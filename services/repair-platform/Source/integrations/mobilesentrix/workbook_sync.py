from __future__ import annotations

import json
import os
import shutil
import time
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, cast

from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

from integrations.mobilesentrix.client import (
    MobileSentrixApiError,
    MobileSentrixClient,
)
from integrations.mobilesentrix.models import (
    MobileSentrixProduct,
)

SHEET_NAME = "40 - Parts"
TABLE_NAME = "tblNormMobileSentrix"

SUPPLIER_NAME = "MobileSentrix"
SUPPLIER_ID = "PSL000001"

SOURCE_NAME = "Mobile Sentrix API"
CHECKPOINT_FILENAME = "mobilesentrix_sync_checkpoint.json"
MAX_RETRY_ATTEMPTS = 5


class MobileSentrixWorkbookSyncError(RuntimeError):
    """Raised when Mobile Sentrix workbook synchronization fails."""


@dataclass(frozen=True, slots=True)
class WorkbookProductMatch:
    row_number: int
    supplier_sku: str
    product_name: str
    unit_cost: Decimal | None
    availability: str
    product_url: str | None


@dataclass(frozen=True, slots=True)
class ProductChange:
    supplier_sku: str
    row_number: int | None
    is_new: bool

    cost_changed: bool
    availability_changed: bool
    name_changed: bool
    url_changed: bool

    existing_cost: Decimal | None
    api_cost: Decimal | None

    existing_availability: str | None
    api_availability: str | None

    existing_name: str | None
    api_name: str | None

    existing_url: str | None
    api_url: str | None


class MobileSentrixWorkbookSyncService:
    """
    Synchronizes Mobile Sentrix supplier observations with the
    Nocturnix master catalog workbook.

    The current implementation intentionally treats the workbook as
    the canonical catalog/pricing source.

    Mobile Sentrix API data may refresh supplier-observation fields,
    but it must not automatically approve products for pricing or
    overwrite Nocturnix classification/governance fields.
    """

    def __init__(
        self,
        *,
        client: MobileSentrixClient | None = None,
        workbook_path: str | Path | None = None,
    ) -> None:
        self.client = client or MobileSentrixClient()

        if workbook_path is None:
            self.workbook_path = self._configured_workbook_path()
        else:
            self.workbook_path = Path(workbook_path).expanduser().resolve()

            if not self.workbook_path.exists():
                raise MobileSentrixWorkbookSyncError(
                    "Workbook does not exist: " f"{self.workbook_path}"
                )

    # ---------------------------------------------------------
    # Configuration
    # ---------------------------------------------------------

    @staticmethod
    def _configured_workbook_path() -> Path:
        """
        Resolve the configured master catalog workbook.

        The dotenv file is loaded explicitly so this module also works
        correctly when imported directly rather than through the
        integrations.mobilesentrix package.
        """

        project_root = Path(__file__).resolve().parents[3]

        env_path = project_root / ".env"

        load_dotenv(
            dotenv_path=env_path,
            override=True,
        )

        configured = os.getenv("NOCTURNIX_MASTER_CATALOG_WORKBOOK")

        if not configured:
            raise MobileSentrixWorkbookSyncError(
                "NOCTURNIX_MASTER_CATALOG_WORKBOOK " "is not configured."
            )

        workbook_path = Path(configured).expanduser().resolve()

        if not workbook_path.exists():
            raise MobileSentrixWorkbookSyncError(
                "Configured workbook does not exist: " f"{workbook_path}"
            )

        return workbook_path

    # ---------------------------------------------------------
    # Dry-run search synchronization
    # ---------------------------------------------------------

    def dry_run_search(
        self,
        *,
        query: str,
        max_results: int = 100,
        start_index: int = 0,
    ) -> dict[str, object]:
        """
        Compare a Mobile Sentrix product search with the workbook.

        This method never modifies the workbook.
        """

        if not query.strip():
            raise ValueError("query must not be empty.")

        existing_products = self._load_existing_products()

        result = self.client.search_products(
            query=query,
            max_results=max_results,
            start_index=start_index,
        )

        data = result.get("data") or {}

        if not isinstance(data, dict):
            raise MobileSentrixWorkbookSyncError(
                "Mobile Sentrix returned an unexpected " "data structure."
            )

        raw_items = data.get("items") or []

        if not isinstance(raw_items, list):
            raise MobileSentrixWorkbookSyncError(
                "Mobile Sentrix product items were not " "returned as a list."
            )

        changes: list[ProductChange] = []

        existing_sku_matches = 0
        new_products = 0

        cost_changes = 0
        availability_changes = 0
        name_changes = 0
        url_changes = 0

        products_with_changes = 0

        for raw_item in raw_items:
            if not isinstance(
                raw_item,
                dict,
            ):
                continue

            product = MobileSentrixProduct.from_api_item(raw_item)

            if not product.supplier_sku:
                continue

            sku = self._normalize_sku(product.supplier_sku)

            if not sku:
                continue

            existing = existing_products.get(sku)

            api_availability = self._availability_from_product(product)

            if existing is None:
                new_products += 1

                change = ProductChange(
                    supplier_sku=sku,
                    row_number=None,
                    is_new=True,
                    cost_changed=False,
                    availability_changed=False,
                    name_changed=False,
                    url_changed=False,
                    existing_cost=None,
                    api_cost=product.unit_cost,
                    existing_availability=None,
                    api_availability=(api_availability),
                    existing_name=None,
                    api_name=product.name,
                    existing_url=None,
                    api_url=product.product_url,
                )

                changes.append(change)
                continue

            existing_sku_matches += 1

            cost_changed = not (
                self._decimal_equal(
                    existing.unit_cost,
                    product.unit_cost,
                )
            )

            availability_changed = (
                self._normalized_text(existing.availability).casefold()
                != self._normalized_text(api_availability).casefold()
            )

            name_changed = self._normalized_text(
                existing.product_name
            ) != self._normalized_text(product.name)

            url_changed = self._normalized_text(
                existing.product_url
            ) != self._normalized_text(product.product_url)

            if cost_changed:
                cost_changes += 1

            if availability_changed:
                availability_changes += 1

            if name_changed:
                name_changes += 1

            if url_changed:
                url_changes += 1

            has_changes = any(
                (
                    cost_changed,
                    availability_changed,
                    name_changed,
                    url_changed,
                )
            )

            if has_changes:
                products_with_changes += 1

            changes.append(
                ProductChange(
                    supplier_sku=sku,
                    row_number=(existing.row_number),
                    is_new=False,
                    cost_changed=cost_changed,
                    availability_changed=(availability_changed),
                    name_changed=name_changed,
                    url_changed=url_changed,
                    existing_cost=(existing.unit_cost),
                    api_cost=product.unit_cost,
                    existing_availability=(existing.availability),
                    api_availability=(api_availability),
                    existing_name=(existing.product_name),
                    api_name=product.name,
                    existing_url=(existing.product_url),
                    api_url=product.product_url,
                )
            )

        return {
            "mode": "dry_run",
            "workbook": str(self.workbook_path),
            "sheet": SHEET_NAME,
            "table": TABLE_NAME,
            "supplier": SUPPLIER_NAME,
            "supplier_id": SUPPLIER_ID,
            "query": query,
            "observed_at": datetime.now(UTC).isoformat(),
            "api_total_items": data.get("total_items"),
            "api_products_returned": len(raw_items),
            "existing_mobile_sentrix_rows": len(existing_products),
            "existing_sku_matches": (existing_sku_matches),
            "new_products": new_products,
            "products_with_changes": (products_with_changes),
            "cost_changes": cost_changes,
            # Retained for compatibility with the existing
            # test/console output. Mobile Sentrix quantity
            # is not a literal inventory count.
            "stock_changes": 0,
            "availability_changes": (availability_changes),
            "name_changes": name_changes,
            "url_changes": url_changes,
            "rows_written": 0,
            "workbook_modified": False,
            "changes": [self._change_to_dict(change) for change in changes],
        }

    # ---------------------------------------------------------
    # Search-based write
    # ---------------------------------------------------------

    def apply_search(
        self,
        *,
        query: str,
        max_results: int = 100,
        start_index: int = 0,
    ) -> dict[str, object]:
        """
        Update existing workbook products returned by a Mobile Sentrix
        search.

        Existing SKU matches only.

        This method intentionally does NOT:
        - append new products;
        - modify Stock Quantity;
        - change verification status;
        - change preferred-cost status;
        - change match status;
        - change selected-for-pricing status;
        - change Nocturnix classification fields.
        """

        if not query.strip():
            raise ValueError("query must not be empty.")

        keep_vba = self.workbook_path.suffix.lower() == ".xlsm"

        workbook = load_workbook(
            filename=self.workbook_path,
            read_only=False,
            data_only=False,
            keep_vba=keep_vba,
        )

        try:
            if SHEET_NAME not in workbook.sheetnames:
                raise MobileSentrixWorkbookSyncError(
                    f"Worksheet not found: {SHEET_NAME}"
                )

            worksheet = workbook[SHEET_NAME]

            if TABLE_NAME not in worksheet.tables:
                raise MobileSentrixWorkbookSyncError(
                    f"Excel table not found: {TABLE_NAME}"
                )

            table = worksheet.tables[TABLE_NAME]

            headers = self._header_map(
                worksheet=worksheet,
                table=table,
            )

            required_headers = {
                "Supplier",
                "part_supplier_id",
                "Supplier SKU",
                "Product Name",
                "Unit Cost",
                "Availability",
                "Product URL",
                "Observed Date",
                "Source File",
            }

            missing_headers = required_headers - set(headers)

            if missing_headers:
                raise MobileSentrixWorkbookSyncError(
                    "Required write columns are missing: "
                    + ", ".join(sorted(missing_headers))
                )

            existing_products = self._load_existing_products()

            result = self.client.search_products(
                query=query,
                max_results=max_results,
                start_index=start_index,
            )

            data = result.get("data") or {}

            if not isinstance(data, dict):
                raise MobileSentrixWorkbookSyncError(
                    "Mobile Sentrix returned an " "unexpected data structure."
                )

            raw_items = data.get("items") or []

            if not isinstance(
                raw_items,
                list,
            ):
                raise MobileSentrixWorkbookSyncError(
                    "Mobile Sentrix product items " "were not returned as a list."
                )

            observed_at = datetime.now(UTC).isoformat()

            rows_updated = 0
            unchanged_rows = 0
            new_products_skipped = 0

            updated_skus: list[str] = []
            skipped_new_skus: list[str] = []

            for raw_item in raw_items:
                if not isinstance(
                    raw_item,
                    dict,
                ):
                    continue

                product = MobileSentrixProduct.from_api_item(raw_item)

                if not product.supplier_sku:
                    continue

                sku = self._normalize_sku(product.supplier_sku)

                if not sku:
                    continue

                existing = existing_products.get(sku)

                if existing is None:
                    new_products_skipped += 1

                    skipped_new_skus.append(sku)

                    continue

                row_number = existing.row_number

                api_availability = self._availability_from_product(product)

                changes_made = False

                live_values: dict[
                    str,
                    object,
                ] = {
                    "Product Name": (product.name),
                    "Unit Cost": (
                        float(product.unit_cost)
                        if product.unit_cost is not None
                        else None
                    ),
                    "Availability": (api_availability),
                    "Product URL": (product.product_url),
                }

                for (
                    column_name,
                    new_value,
                ) in live_values.items():

                    cell = self._writable_cell(
                        worksheet=worksheet,
                        row=row_number,
                        column=headers[column_name],
                    )

                    if column_name == "Unit Cost":
                        old_decimal = self._to_decimal(cell.value)

                        if self._decimal_equal(
                            old_decimal,
                            product.unit_cost,
                        ):
                            continue

                    elif self._normalized_text(cell.value) == self._normalized_text(
                        new_value
                    ):
                        continue

                    self._set_cell_value(
                        cell=cell,
                        value=new_value,
                    )
                    changes_made = True

                self._writable_cell(
                    worksheet=worksheet,
                    row=row_number,
                    column=headers["Observed Date"],
                ).value = observed_at

                self._writable_cell(
                    worksheet=worksheet,
                    row=row_number,
                    column=headers["Source File"],
                ).value = SOURCE_NAME

                if changes_made:
                    rows_updated += 1

                    updated_skus.append(sku)
                else:
                    unchanged_rows += 1

            backup_path = self._create_backup_path(label=("before-mobilesentrix-sync"))

            shutil.copy2(
                self.workbook_path,
                backup_path,
            )

            workbook.save(self.workbook_path)

            return {
                "mode": "write_existing_only",
                "workbook": str(self.workbook_path),
                "backup": str(backup_path),
                "sheet": SHEET_NAME,
                "table": TABLE_NAME,
                "supplier": SUPPLIER_NAME,
                "query": query,
                "api_total_items": data.get("total_items"),
                "api_products_returned": len(raw_items),
                "rows_updated": rows_updated,
                "unchanged_rows": (unchanged_rows),
                "new_products_skipped": (new_products_skipped),
                "updated_skus": (updated_skus),
                "skipped_new_skus": (skipped_new_skus),
                "workbook_modified": True,
            }

        except PermissionError as exc:
            raise MobileSentrixWorkbookSyncError(
                "The workbook could not be saved. " "Make sure it is not open in Excel."
            ) from exc

        finally:
            workbook.close()

    # ---------------------------------------------------------
    # Existing-product batch synchronization
    # ---------------------------------------------------------

    def sync_existing_batch(
        self,
        *,
        start_row: int | None = None,
        batch_size: int = 50,
        request_delay_seconds: float = 0.25,
    ) -> dict[str, object]:
        """
        Refresh a controlled batch of existing Mobile Sentrix products.

        Each workbook Supplier SKU is searched individually.

        A result is accepted only when product_code exactly matches
        the workbook Supplier SKU.

        Mobile Sentrix quantity is used ONLY to derive Availability.
        It is not written to the workbook Stock Quantity field.
        """

        if batch_size < 1:
            raise ValueError("batch_size must be at least 1.")

        if request_delay_seconds < 0:
            raise ValueError("request_delay_seconds cannot " "be negative.")

        keep_vba = self.workbook_path.suffix.lower() == ".xlsm"

        workbook = load_workbook(
            filename=self.workbook_path,
            read_only=False,
            data_only=False,
            keep_vba=keep_vba,
        )

        try:
            if SHEET_NAME not in workbook.sheetnames:
                raise MobileSentrixWorkbookSyncError(
                    f"Worksheet not found: {SHEET_NAME}"
                )

            worksheet = workbook[SHEET_NAME]

            if TABLE_NAME not in worksheet.tables:
                raise MobileSentrixWorkbookSyncError(
                    f"Excel table not found: {TABLE_NAME}"
                )

            table = worksheet.tables[TABLE_NAME]

            headers = self._header_map(
                worksheet=worksheet,
                table=table,
            )

            required_headers = {
                "Supplier",
                "part_supplier_id",
                "Supplier SKU",
                "Product Name",
                "Unit Cost",
                "Availability",
                "Product URL",
                "Observed Date",
                "Source File",
            }

            missing_headers = required_headers - set(headers)

            if missing_headers:
                raise MobileSentrixWorkbookSyncError(
                    "Required sync columns are missing: "
                    + ", ".join(sorted(missing_headers))
                )

            (
                _,
                min_row,
                _,
                max_row,
            ) = self._table_bounds(table)

            first_data_row = min_row + 1

            if start_row is None:
                scan_start = first_data_row
            else:
                scan_start = max(
                    start_row,
                    first_data_row,
                )

            candidate_rows: list[tuple[int, str]] = []

            for row_number in range(
                scan_start,
                max_row + 1,
            ):
                supplier = self._normalized_text(
                    worksheet.cell(
                        row=row_number,
                        column=headers["Supplier"],
                    ).value
                )

                supplier_id = self._normalized_text(
                    worksheet.cell(
                        row=row_number,
                        column=headers["part_supplier_id"],
                    ).value
                )

                supplier_matches = supplier.casefold() == SUPPLIER_NAME.casefold()

                supplier_id_matches = supplier_id == SUPPLIER_ID

                if not (supplier_matches or supplier_id_matches):
                    continue

                sku = self._normalize_sku(
                    worksheet.cell(
                        row=row_number,
                        column=headers["Supplier SKU"],
                    ).value
                )

                if not sku:
                    continue

                candidate_rows.append(
                    (
                        row_number,
                        sku,
                    )
                )

                if len(candidate_rows) >= batch_size:
                    break

            if not candidate_rows:
                return {
                    "mode": "batch_sync",
                    "workbook": str(self.workbook_path),
                    "batch_size_requested": (batch_size),
                    "products_processed": 0,
                    "exact_matches": 0,
                    "rows_updated": 0,
                    "unchanged_rows": 0,
                    "not_found": 0,
                    "ambiguous_matches": 0,
                    "api_errors": 0,
                    "next_start_row": None,
                    "workbook_modified": False,
                    "message": ("No additional Mobile Sentrix " "rows were found."),
                    "errors": [],
                }

            backup_path = self._create_backup_path(
                label=("before-mobilesentrix-batch-sync")
            )

            shutil.copy2(
                self.workbook_path,
                backup_path,
            )

            processed = 0
            exact_matches = 0
            rows_updated = 0
            unchanged_rows = 0
            not_found = 0
            ambiguous_matches = 0
            api_errors = 0

            errors: list[dict[str, object]] = []

            updated_skus: list[str] = []

            observed_at = datetime.now(UTC).isoformat()

            for (
                row_number,
                sku,
            ) in candidate_rows:

                processed += 1

                try:
                    result = self.client.search_products(
                        query=sku,
                        max_results=10,
                        start_index=0,
                    )

                    data = result.get("data") or {}

                    if not isinstance(
                        data,
                        dict,
                    ):
                        raise (
                            MobileSentrixWorkbookSyncError(
                                "Unexpected API " "data structure."
                            )
                        )

                    raw_items = data.get("items") or []

                    if not isinstance(
                        raw_items,
                        list,
                    ):
                        raise (
                            MobileSentrixWorkbookSyncError(
                                "Unexpected API " "items structure."
                            )
                        )

                    exact_items: list[dict[str, Any]] = []

                    for item in raw_items:
                        if not isinstance(
                            item,
                            dict,
                        ):
                            continue

                        api_sku = self._normalize_sku(item.get("product_code"))

                        if api_sku == sku:
                            exact_items.append(item)

                    if not exact_items:
                        not_found += 1

                        errors.append(
                            {
                                "row": (row_number),
                                "sku": sku,
                                "error": (
                                    "No exact " "product_code " "match returned."
                                ),
                            }
                        )

                        continue

                    if len(exact_items) > 1:
                        ambiguous_matches += 1

                        errors.append(
                            {
                                "row": (row_number),
                                "sku": sku,
                                "error": (
                                    "Multiple exact " "SKU matches " "were returned."
                                ),
                            }
                        )

                        continue

                    exact_matches += 1

                    product = MobileSentrixProduct.from_api_item(exact_items[0])

                    api_availability = self._availability_from_product(product)

                    changes_made = False

                    live_values: dict[
                        str,
                        object,
                    ] = {
                        "Product Name": (product.name),
                        "Unit Cost": (
                            float(product.unit_cost)
                            if product.unit_cost is not None
                            else None
                        ),
                        "Availability": (api_availability),
                        "Product URL": (product.product_url),
                    }

                    for (
                        column_name,
                        new_value,
                    ) in live_values.items():

                        cell = self._writable_cell(
                            worksheet=worksheet,
                            row=row_number,
                            column=headers[column_name],
                        )

                        if column_name == "Unit Cost":
                            old_decimal = self._to_decimal(cell.value)

                            if self._decimal_equal(
                                old_decimal,
                                product.unit_cost,
                            ):
                                continue

                        elif self._normalized_text(cell.value) == self._normalized_text(
                            new_value
                        ):
                            continue

                        self._set_cell_value(
                            cell=cell,
                            value=new_value,
                        )
                        changes_made = True

                    self._writable_cell(
                        worksheet=worksheet,
                        row=row_number,
                        column=headers["Observed Date"],
                    ).value = observed_at

                    self._writable_cell(
                        worksheet=worksheet,
                        row=row_number,
                        column=headers["Source File"],
                    ).value = SOURCE_NAME

                    if changes_made:
                        rows_updated += 1

                        updated_skus.append(sku)
                    else:
                        unchanged_rows += 1

                except (
                    MobileSentrixApiError,
                    MobileSentrixWorkbookSyncError,
                    ValueError,
                    TypeError,
                ) as exc:
                    api_errors += 1

                    errors.append(
                        {
                            "row": row_number,
                            "sku": sku,
                            "error": (f"{type(exc).__name__}: " f"{exc}"),
                        }
                    )

                if request_delay_seconds > 0:
                    time.sleep(request_delay_seconds)

            workbook.save(self.workbook_path)

            last_processed_row = candidate_rows[-1][0]

            if last_processed_row < max_row:
                next_start_row = last_processed_row + 1
            else:
                next_start_row = None

            return {
                "mode": "batch_sync",
                "workbook": str(self.workbook_path),
                "backup": str(backup_path),
                "batch_size_requested": (batch_size),
                "products_processed": (processed),
                "exact_matches": (exact_matches),
                "rows_updated": (rows_updated),
                "unchanged_rows": (unchanged_rows),
                "not_found": not_found,
                "ambiguous_matches": (ambiguous_matches),
                "api_errors": api_errors,
                "next_start_row": (next_start_row),
                "updated_skus": (updated_skus),
                "errors": errors,
                "workbook_modified": True,
            }

        except PermissionError as exc:
            raise MobileSentrixWorkbookSyncError(
                "The workbook could not be saved. " "Make sure it is not open in Excel."
            ) from exc

        finally:
            workbook.close()

    def sync_next_batch(
        self,
        *,
        batch_size: int = 50,
        request_delay_seconds: float = 0.25,
    ) -> dict[str, object]:
        """
        Synchronize the next workbook batch using the persistent
        checkpoint.

        Failed, missing, or ambiguous exact-SKU lookups are retained
        in the retry queue before primary scan progress advances.
        """

        checkpoint = self._load_checkpoint()
        retry_items = self._retry_items_from_checkpoint(checkpoint)

        if bool(checkpoint.get("cycle_complete", False)):
            return {
                "mode": "checkpoint_batch_sync",
                "cycle_complete": True,
                "products_processed": 0,
                "retry_count": len(retry_items),
                "message": "The current Mobile Sentrix synchronization cycle is already complete.",
                "checkpoint_path": str(self._checkpoint_path()),
                "checkpoint": checkpoint,
            }

        next_start_row_value = checkpoint.get("next_start_row")

        # Primary scan has reached the end, but unresolved retry items
        # still exist. Do not restart at row 2.
        if next_start_row_value is None:
            pending_retry_items = [
                item
                for item in retry_items
                if str(item.get("status", "pending")).lower()
                == "pending"
            ]

            if pending_retry_items:
                return {
                    "mode": "checkpoint_batch_sync",
                    "cycle_complete": False,
                    "products_processed": 0,
                    "retry_count": len(retry_items),
                    "message": "Primary scan is complete. Retry items remain and must be processed.",
                    "checkpoint_path": str(self._checkpoint_path()),
                    "checkpoint": checkpoint,
                }

            checkpoint["cycle_complete"] = True
            checkpoint["last_completed_at"] = datetime.now(UTC).isoformat()
            self._save_checkpoint(checkpoint)
            return {
                "mode": "checkpoint_batch_sync",
                "cycle_complete": True,
                "products_processed": 0,
                "retry_count": 0,
                "message": "The Mobile Sentrix synchronization cycle is complete.",
                "checkpoint_path": str(self._checkpoint_path()),
                "checkpoint": checkpoint,
            }

        start_row = self._to_int(next_start_row_value)
        if start_row is None:
            raise MobileSentrixWorkbookSyncError(
                "The Mobile Sentrix checkpoint contains an invalid next_start_row."
            )

        now = datetime.now(UTC).isoformat()
        if not checkpoint.get("cycle_started_at"):
            checkpoint["cycle_started_at"] = now

        result = self.sync_existing_batch(
            start_row=start_row,
            batch_size=batch_size,
            request_delay_seconds=request_delay_seconds,
        )

        processed = self._to_int(result.get("products_processed")) or 0
        previous_processed = (
            self._to_int(checkpoint.get("products_processed_this_cycle")) or 0
        )
        checkpoint["products_processed_this_cycle"] = previous_processed + processed
        checkpoint["last_run_at"] = datetime.now(UTC).isoformat()

        self._capture_batch_errors(
            retry_items=retry_items,
            batch_errors=result.get("errors"),
        )
        checkpoint["retry_items"] = retry_items

        result_next_row = result.get("next_start_row")
        if result_next_row is None:
            checkpoint["next_start_row"] = None
        else:
            parsed_next_row = self._to_int(result_next_row)
            if parsed_next_row is None:
                raise MobileSentrixWorkbookSyncError(
                    "Batch synchronization returned an invalid next_start_row."
                )
            checkpoint["next_start_row"] = parsed_next_row

        primary_scan_complete = checkpoint.get("next_start_row") is None

        pending_retry_count = sum(
            1
            for item in retry_items
            if str(item.get("status", "pending")).lower() == "pending"
        )

        cycle_complete = primary_scan_complete and pending_retry_count == 0
        checkpoint["cycle_complete"] = cycle_complete

        if cycle_complete:
            checkpoint["last_completed_at"] = datetime.now(UTC).isoformat()

        self._save_checkpoint(checkpoint)

        result["mode"] = "checkpoint_batch_sync"
        result["cycle_complete"] = cycle_complete
        result["products_processed_this_cycle"] = checkpoint[
            "products_processed_this_cycle"
        ]
        result["retry_count"] = len(retry_items)
        result["checkpoint_path"] = str(self._checkpoint_path())
        result["checkpoint"] = checkpoint
        return result

    def retry_failed_items(
        self,
        *,
        max_items: int = 25,
        request_delay_seconds: float = 0.25,
    ) -> dict[str, object]:
        """
        Retry genuinely retryable Mobile Sentrix SKU lookups stored in
        the persistent checkpoint.

        Successful exact matches are removed from retry state.

        A successful supplier response that contains no exact
        product_code match is classified as status=unmatched and is
        retained for supplier/manual investigation without consuming
        additional retry attempts.

        Retryable failures remain status=pending until they reach
        MAX_RETRY_ATTEMPTS, at which point they remain recorded with
        status=exhausted.

        Unmatched and exhausted records are terminal review states and
        do not prevent the synchronization cycle from completing.
        """

        if max_items < 1:
            raise ValueError("max_items must be at least 1.")
        if request_delay_seconds < 0:
            raise ValueError("request_delay_seconds cannot be negative.")

        checkpoint = self._load_checkpoint()
        retry_items = self._retry_items_from_checkpoint(checkpoint)

        pending_items = [
            item
            for item in retry_items
            if str(item.get("status", "pending")).lower() == "pending"
        ]

        terminal_items = [
            item
            for item in retry_items
            if str(item.get("status", "pending")).lower() != "pending"
        ]

        if not pending_items:
            primary_scan_complete = checkpoint.get("next_start_row") is None
            checkpoint["cycle_complete"] = primary_scan_complete

            if primary_scan_complete:
                checkpoint["last_completed_at"] = datetime.now(UTC).isoformat()

            self._save_checkpoint(checkpoint)

            unmatched_count = sum(
                1
                for item in retry_items
                if str(item.get("status", "")).lower() == "unmatched"
            )
            exhausted_count = sum(
                1
                for item in retry_items
                if str(item.get("status", "")).lower() == "exhausted"
            )

            return {
                "mode": "retry_failed_items",
                "retry_items_processed": 0,
                "recovered": 0,
                "still_pending": 0,
                "unmatched": unmatched_count,
                "exhausted": exhausted_count,
                "api_errors": 0,
                "rows_updated": 0,
                "unchanged_rows": 0,
                "retry_count": len(retry_items),
                "supplier_data_changed": False,
                "observation_metadata_refreshed": False,
                "workbook_saved": False,
                "workbook_modified": False,
                "cycle_complete": bool(checkpoint.get("cycle_complete", False)),
                "checkpoint_path": str(self._checkpoint_path()),
                "checkpoint": checkpoint,
            }

        keep_vba = self.workbook_path.suffix.lower() == ".xlsm"
        workbook = load_workbook(
            filename=self.workbook_path,
            read_only=False,
            data_only=False,
            keep_vba=keep_vba,
        )

        try:
            if SHEET_NAME not in workbook.sheetnames:
                raise MobileSentrixWorkbookSyncError(
                    f"Worksheet not found: {SHEET_NAME}"
                )

            worksheet = workbook[SHEET_NAME]

            if TABLE_NAME not in worksheet.tables:
                raise MobileSentrixWorkbookSyncError(
                    f"Excel table not found: {TABLE_NAME}"
                )

            table = worksheet.tables[TABLE_NAME]
            headers = self._header_map(
                worksheet=worksheet,
                table=table,
            )

            required_headers = {
                "Supplier SKU",
                "Product Name",
                "Unit Cost",
                "Availability",
                "Product URL",
                "Observed Date",
                "Source File",
            }

            missing_headers = required_headers - set(headers)
            if missing_headers:
                raise MobileSentrixWorkbookSyncError(
                    "Required retry columns are missing: "
                    + ", ".join(sorted(missing_headers))
                )

            selected = pending_items[:max_items]
            untouched_pending = pending_items[max_items:]
            remaining_pending: list[dict[str, object]] = []
            newly_terminal: list[dict[str, object]] = []

            processed = 0
            recovered = 0
            api_errors = 0
            rows_updated = 0
            unchanged_rows = 0
            supplier_data_changed = False
            observation_metadata_refreshed = False
            observed_at = datetime.now(UTC).isoformat()

            for item in selected:
                row_number = self._to_int(item.get("row"))
                sku = self._normalize_sku(item.get("sku"))
                attempts = self._to_int(item.get("attempts")) or 0

                if row_number is None or not sku:
                    item["status"] = "exhausted"
                    item["last_error"] = "Retry item has an invalid row or SKU."
                    newly_terminal.append(item)
                    continue

                if attempts >= MAX_RETRY_ATTEMPTS:
                    item["status"] = "exhausted"
                    newly_terminal.append(item)
                    continue

                processed += 1

                try:
                    result = self.client.search_products(
                        query=sku,
                        max_results=10,
                        start_index=0,
                    )

                    data = result.get("data") or {}
                    if not isinstance(data, dict):
                        raise MobileSentrixWorkbookSyncError(
                            "Unexpected API data structure."
                        )

                    raw_items = data.get("items") or []
                    if not isinstance(raw_items, list):
                        raise MobileSentrixWorkbookSyncError(
                            "Unexpected API items structure."
                        )

                    exact_items: list[dict[str, Any]] = []

                    for raw_item in raw_items:
                        if not isinstance(raw_item, dict):
                            continue

                        if self._normalize_sku(raw_item.get("product_code")) == sku:
                            exact_items.append(raw_item)

                    if not exact_items:
                        timestamp = datetime.now(UTC).isoformat()

                        item["status"] = "unmatched"
                        item["last_error"] = "No exact product_code match returned."
                        item["last_failed_at"] = timestamp

                        if not item.get("first_failed_at"):
                            item["first_failed_at"] = timestamp

                        newly_terminal.append(item)

                    elif len(exact_items) > 1:
                        raise MobileSentrixWorkbookSyncError(
                            "Multiple exact SKU matches were returned."
                        )

                    else:
                        product = MobileSentrixProduct.from_api_item(exact_items[0])

                        api_availability = self._availability_from_product(product)
                        changes_made = False

                        live_values: dict[str, object] = {
                            "Product Name": product.name,
                            "Unit Cost": (
                                float(product.unit_cost)
                                if product.unit_cost is not None
                                else None
                            ),
                            "Availability": api_availability,
                            "Product URL": product.product_url,
                        }

                        for column_name, new_value in live_values.items():
                            cell = self._writable_cell(
                                worksheet=worksheet,
                                row=row_number,
                                column=headers[column_name],
                            )

                            if column_name == "Unit Cost":
                                old_decimal = self._to_decimal(cell.value)

                                if self._decimal_equal(
                                    old_decimal,
                                    product.unit_cost,
                                ):
                                    continue

                            elif self._normalized_text(
                                cell.value
                            ) == self._normalized_text(new_value):
                                continue

                            self._set_cell_value(
                                cell=cell,
                                value=new_value,
                            )

                            changes_made = True
                            supplier_data_changed = True

                        observed_cell = self._writable_cell(
                            worksheet=worksheet,
                            row=row_number,
                            column=headers["Observed Date"],
                        )

                        source_cell = self._writable_cell(
                            worksheet=worksheet,
                            row=row_number,
                            column=headers["Source File"],
                        )

                        observed_cell.value = observed_at
                        source_cell.value = SOURCE_NAME
                        observation_metadata_refreshed = True

                        if changes_made:
                            rows_updated += 1
                        else:
                            unchanged_rows += 1

                        recovered += 1

                except (
                    MobileSentrixApiError,
                    MobileSentrixWorkbookSyncError,
                    ValueError,
                    TypeError,
                ) as exc:
                    api_errors += 1

                    new_attempts = attempts + 1
                    timestamp = datetime.now(UTC).isoformat()

                    item["attempts"] = new_attempts
                    item["last_error"] = f"{type(exc).__name__}: {exc}"
                    item["last_failed_at"] = timestamp

                    if not item.get("first_failed_at"):
                        item["first_failed_at"] = timestamp

                    if new_attempts >= MAX_RETRY_ATTEMPTS:
                        item["status"] = "exhausted"
                        newly_terminal.append(item)
                    else:
                        item["status"] = "pending"
                        remaining_pending.append(item)

                if request_delay_seconds > 0:
                    time.sleep(request_delay_seconds)

            final_retry_items = (
                remaining_pending + untouched_pending + terminal_items + newly_terminal
            )

            checkpoint["retry_items"] = final_retry_items
            checkpoint["last_retry_run_at"] = datetime.now(UTC).isoformat()

            workbook_needs_save = (
                supplier_data_changed or observation_metadata_refreshed
            )

            backup_path: Path | None = None

            if workbook_needs_save:
                backup_path = self._create_backup_path(
                    label="before-mobilesentrix-retry-sync"
                )

                shutil.copy2(
                    self.workbook_path,
                    backup_path,
                )

                workbook.save(self.workbook_path)

            pending_count = sum(
                1
                for retry_item in final_retry_items
                if str(retry_item.get("status", "pending")).lower() == "pending"
            )

            unmatched_count = sum(
                1
                for retry_item in final_retry_items
                if str(retry_item.get("status", "")).lower() == "unmatched"
            )

            exhausted_count = sum(
                1
                for retry_item in final_retry_items
                if str(retry_item.get("status", "")).lower() == "exhausted"
            )

            primary_scan_complete = checkpoint.get("next_start_row") is None

            cycle_complete = primary_scan_complete and pending_count == 0

            checkpoint["cycle_complete"] = cycle_complete

            if cycle_complete:
                checkpoint["last_completed_at"] = datetime.now(UTC).isoformat()

            self._save_checkpoint(checkpoint)

            return {
                "mode": "retry_failed_items",
                "retry_items_processed": processed,
                "recovered": recovered,
                "still_pending": pending_count,
                "unmatched": unmatched_count,
                "exhausted": exhausted_count,
                "api_errors": api_errors,
                "rows_updated": rows_updated,
                "unchanged_rows": unchanged_rows,
                "retry_count": len(final_retry_items),
                "supplier_data_changed": supplier_data_changed,
                "observation_metadata_refreshed": (observation_metadata_refreshed),
                "workbook_saved": workbook_needs_save,
                "workbook_modified": workbook_needs_save,
                "backup": (str(backup_path) if backup_path is not None else None),
                "cycle_complete": cycle_complete,
                "checkpoint_path": str(self._checkpoint_path()),
                "checkpoint": checkpoint,
            }

        except PermissionError as exc:
            raise MobileSentrixWorkbookSyncError(
                "The workbook could not be saved. " "Make sure it is not open in Excel."
            ) from exc

        finally:
            workbook.close()

    def reset_sync_checkpoint(
        self,
    ) -> dict[str, object]:
        """Start a new synchronization cycle and clear retry state."""

        checkpoint: dict[str, object] = {
            "next_start_row": 2,
            "cycle_started_at": None,
            "last_run_at": None,
            "last_retry_run_at": None,
            "last_completed_at": None,
            "products_processed_this_cycle": 0,
            "cycle_complete": False,
            "retry_items": [],
        }
        self._save_checkpoint(checkpoint)
        return checkpoint

    def get_sync_checkpoint(
        self,
    ) -> dict[str, object]:
        """Return the current synchronization checkpoint."""

        return self._load_checkpoint()

    def _checkpoint_path(self) -> Path:
        """Return the persistent Mobile Sentrix checkpoint path."""

        project_root = Path(__file__).resolve().parents[3]
        data_folder = project_root / "Data"
        data_folder.mkdir(parents=True, exist_ok=True)
        return data_folder / CHECKPOINT_FILENAME

    def _load_checkpoint(
        self,
    ) -> dict[str, object]:
        """Load and schema-normalize the persistent checkpoint."""

        checkpoint_path = self._checkpoint_path()
        if not checkpoint_path.exists():
            return self.reset_sync_checkpoint()

        try:
            with checkpoint_path.open("r", encoding="utf-8") as file:
                loaded = json.load(file)
        except (OSError, json.JSONDecodeError) as exc:
            raise MobileSentrixWorkbookSyncError(
                f"Unable to read Mobile Sentrix checkpoint: {checkpoint_path}"
            ) from exc

        if not isinstance(loaded, dict):
            raise MobileSentrixWorkbookSyncError(
                "Mobile Sentrix checkpoint must contain a JSON object."
            )

        data: dict[str, object] = {str(key): value for key, value in loaded.items()}
        defaults: dict[str, object] = {
            "next_start_row": 2,
            "cycle_started_at": None,
            "last_run_at": None,
            "last_retry_run_at": None,
            "last_completed_at": None,
            "products_processed_this_cycle": 0,
            "cycle_complete": False,
            "retry_items": [],
        }
        changed = False
        for key, default_value in defaults.items():
            if key not in data:
                data[key] = default_value
                changed = True

        if not isinstance(data.get("retry_items"), list):
            raise MobileSentrixWorkbookSyncError(
                "Mobile Sentrix checkpoint retry_items must be a JSON list."
            )

        if changed:
            self._save_checkpoint(data)
        return data

    def _save_checkpoint(
        self,
        checkpoint: dict[str, object],
    ) -> None:
        """Atomically save the synchronization checkpoint."""

        checkpoint_path = self._checkpoint_path()
        temporary_path = checkpoint_path.with_suffix(checkpoint_path.suffix + ".tmp")
        try:
            with temporary_path.open("w", encoding="utf-8") as file:
                json.dump(checkpoint, file, indent=2, sort_keys=True)
                file.write("\n")
            os.replace(temporary_path, checkpoint_path)
        except OSError as exc:
            try:
                if temporary_path.exists():
                    temporary_path.unlink()
            except OSError:
                pass
            raise MobileSentrixWorkbookSyncError(
                f"Unable to save Mobile Sentrix checkpoint: {checkpoint_path}"
            ) from exc

    def _retry_items_from_checkpoint(
        self,
        checkpoint: dict[str, object],
    ) -> list[dict[str, object]]:
        """Return a normalized mutable retry-item list."""

        raw_retry_items = checkpoint.get("retry_items")
        if raw_retry_items is None:
            return []
        if not isinstance(raw_retry_items, list):
            raise MobileSentrixWorkbookSyncError(
                "Mobile Sentrix checkpoint retry_items must be a list."
            )

        retry_items: list[dict[str, object]] = []
        for raw_item in raw_retry_items:
            if not isinstance(raw_item, dict):
                continue
            retry_items.append({str(key): value for key, value in raw_item.items()})
        return retry_items

    def _capture_batch_errors(
        self,
        *,
        retry_items: list[dict[str, object]],
        batch_errors: object,
    ) -> None:
        """
        Add or update failed batch items in persistent retry state.

        A successful Mobile Sentrix lookup that returned no exact
        product_code match is classified as unmatched immediately.
        Other failures remain retryable until MAX_RETRY_ATTEMPTS.
        """

        if not isinstance(batch_errors, list):
            return

        existing: dict[
            tuple[int, str],
            dict[str, object],
        ] = {}

        for item in retry_items:
            row_number = self._to_int(item.get("row"))
            sku = self._normalize_sku(item.get("sku"))

            if row_number is not None and sku:
                existing[(row_number, sku)] = item

        for raw_error in batch_errors:
            if not isinstance(raw_error, dict):
                continue

            error = {str(key): value for key, value in raw_error.items()}

            row_number = self._to_int(error.get("row"))
            sku = self._normalize_sku(error.get("sku"))

            if row_number is None or not sku:
                continue

            timestamp = datetime.now(UTC).isoformat()
            error_text = str(error.get("error", ""))

            is_unmatched = "No exact product_code match returned." in error_text

            key = (row_number, sku)
            existing_item = existing.get(key)

            if existing_item is not None:
                existing_item["last_error"] = error_text
                existing_item["last_failed_at"] = timestamp

                if not existing_item.get("first_failed_at"):
                    existing_item["first_failed_at"] = timestamp

                if is_unmatched:
                    existing_item["status"] = "unmatched"
                    continue

                attempts = self._to_int(existing_item.get("attempts")) or 0

                new_attempts = attempts + 1
                existing_item["attempts"] = new_attempts

                existing_item["status"] = (
                    "exhausted" if new_attempts >= MAX_RETRY_ATTEMPTS else "pending"
                )

                continue

            new_item: dict[str, object] = {
                "row": row_number,
                "sku": sku,
                "last_error": error_text,
                "attempts": 1,
                "status": ("unmatched" if is_unmatched else "pending"),
                "first_failed_at": timestamp,
                "last_failed_at": timestamp,
            }

            retry_items.append(new_item)
            existing[key] = new_item

    # ---------------------------------------------------------
    # Workbook loading
    # ---------------------------------------------------------

    def _load_existing_products(
        self,
    ) -> dict[str, WorkbookProductMatch]:
        """
        Load existing Mobile Sentrix products indexed by Supplier SKU.
        """

        keep_vba = self.workbook_path.suffix.lower() == ".xlsm"

        workbook = load_workbook(
            filename=self.workbook_path,
            read_only=False,
            data_only=False,
            keep_vba=keep_vba,
        )

        try:
            if SHEET_NAME not in workbook.sheetnames:
                raise MobileSentrixWorkbookSyncError(
                    f"Worksheet not found: {SHEET_NAME}"
                )

            worksheet = workbook[SHEET_NAME]

            if TABLE_NAME not in worksheet.tables:
                raise MobileSentrixWorkbookSyncError(
                    f"Excel table not found: {TABLE_NAME}"
                )

            table = worksheet.tables[TABLE_NAME]

            headers = self._header_map(
                worksheet=worksheet,
                table=table,
            )

            required_headers = {
                "Supplier",
                "part_supplier_id",
                "Supplier SKU",
                "Product Name",
                "Unit Cost",
                "Availability",
                "Product URL",
            }

            missing_headers = required_headers - set(headers)

            if missing_headers:
                raise MobileSentrixWorkbookSyncError(
                    "Required workbook columns "
                    "are missing: " + ", ".join(sorted(missing_headers))
                )

            (
                _,
                min_row,
                _,
                max_row,
            ) = self._table_bounds(table)

            products: dict[
                str,
                WorkbookProductMatch,
            ] = {}

            for row_number in range(
                min_row + 1,
                max_row + 1,
            ):
                supplier = self._normalized_text(
                    worksheet.cell(
                        row=row_number,
                        column=headers["Supplier"],
                    ).value
                )

                supplier_id = self._normalized_text(
                    worksheet.cell(
                        row=row_number,
                        column=headers["part_supplier_id"],
                    ).value
                )

                supplier_matches = supplier.casefold() == SUPPLIER_NAME.casefold()

                supplier_id_matches = supplier_id == SUPPLIER_ID

                if not (supplier_matches or supplier_id_matches):
                    continue

                sku = self._normalize_sku(
                    worksheet.cell(
                        row=row_number,
                        column=headers["Supplier SKU"],
                    ).value
                )

                if not sku:
                    continue

                products[sku] = WorkbookProductMatch(
                    row_number=row_number,
                    supplier_sku=sku,
                    product_name=(
                        self._normalized_text(
                            worksheet.cell(
                                row=row_number,
                                column=headers["Product Name"],
                            ).value
                        )
                    ),
                    unit_cost=(
                        self._to_decimal(
                            worksheet.cell(
                                row=row_number,
                                column=headers["Unit Cost"],
                            ).value
                        )
                    ),
                    availability=(
                        self._normalized_text(
                            worksheet.cell(
                                row=row_number,
                                column=headers["Availability"],
                            ).value
                        )
                    ),
                    product_url=(
                        self._normalized_optional_text(
                            worksheet.cell(
                                row=row_number,
                                column=headers["Product URL"],
                            ).value
                        )
                    ),
                )

            return products

        finally:
            workbook.close()

    # ---------------------------------------------------------
    # Table helpers
    # ---------------------------------------------------------

    @classmethod
    def _header_map(
        cls,
        *,
        worksheet: Worksheet,
        table: Table,
    ) -> dict[str, int]:
        """
        Return table header name -> worksheet column number.
        """

        (
            min_col,
            min_row,
            max_col,
            _,
        ) = cls._table_bounds(table)

        headers: dict[
            str,
            int,
        ] = {}

        for column_number in range(
            min_col,
            max_col + 1,
        ):
            value = worksheet.cell(
                row=min_row,
                column=column_number,
            ).value

            header = cls._normalized_text(value)

            if header:
                headers[header] = column_number

        return headers

    @staticmethod
    def _table_bounds(
        table: Table,
    ) -> tuple[int, int, int, int]:
        """Convert an Excel table reference into numeric bounds."""

        from openpyxl.utils.cell import range_boundaries

        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if min_col is None or min_row is None or max_col is None or max_row is None:
            raise MobileSentrixWorkbookSyncError(
                "Excel table reference does not contain complete bounds: "
                f"{table.ref}"
            )
        return min_col, min_row, max_col, max_row

    @staticmethod
    def _set_cell_value(
        *,
        cell: Cell,
        value: object,
    ) -> None:
        """Assign a supported scalar value to a writable Excel cell."""

        if value is None:
            cell.value = None
            return

        if isinstance(value, bool):
            cell.value = value
            return

        if isinstance(value, str):
            cell.value = value
            return

        if isinstance(value, int):
            cell.value = value
            return

        if isinstance(value, float):
            cell.value = value
            return

        if isinstance(value, Decimal):
            cell.value = value
            return

        if isinstance(value, datetime):
            cell.value = value
            return

        if isinstance(value, date):
            cell.value = value
            return

        raise MobileSentrixWorkbookSyncError(
            "Unsupported Excel cell value type: " f"{type(value).__name__}"
        )

    @staticmethod
    def _writable_cell(
        *,
        worksheet: Worksheet,
        row: int,
        column: int,
    ) -> Cell:
        """Return a normal writable openpyxl Cell."""

        cell = worksheet.cell(row=row, column=column)
        if not isinstance(cell, Cell):
            raise MobileSentrixWorkbookSyncError(
                f"Expected writable cell at row {row}, column {column}."
            )
        return cast(Cell, cell)

    # ---------------------------------------------------------
    # Mobile Sentrix interpretation
    # ---------------------------------------------------------

    @staticmethod
    def _availability_from_product(
        product: MobileSentrixProduct,
    ) -> str:
        """
        Translate Mobile Sentrix's binary stock indicator into a
        Nocturnix availability value.

        Mobile Sentrix documentation defines:
            quantity = 1 -> In Stock
            quantity = 0 -> Out of Stock
        """

        if product.in_stock is True:
            return "In Stock"

        if product.in_stock is False:
            return "Out of Stock"

        return "Unknown"

    # ---------------------------------------------------------
    # Normalization helpers
    # ---------------------------------------------------------

    @staticmethod
    def _normalize_sku(
        value: object,
    ) -> str:
        if value is None:
            return ""

        if isinstance(
            value,
            float,
        ):
            if value.is_integer():
                return str(int(value))

        if isinstance(
            value,
            Decimal,
        ):
            if value == value.to_integral_value():
                return str(int(value))

        text = str(value).strip()

        if text.endswith(".0"):
            integer_part = text[:-2]

            if integer_part.isdigit():
                return integer_part

        return text

    @staticmethod
    def _normalized_text(
        value: object,
    ) -> str:
        if value is None:
            return ""

        return str(value).strip()

    @classmethod
    def _normalized_optional_text(
        cls,
        value: object,
    ) -> str | None:
        text = cls._normalized_text(value)

        return text or None

    @staticmethod
    def _to_decimal(
        value: object,
    ) -> Decimal | None:
        if value is None:
            return None

        if isinstance(
            value,
            Decimal,
        ):
            return value

        text = str(value).strip()

        if not text:
            return None

        text = text.replace(
            "$",
            "",
        ).replace(
            ",",
            "",
        )

        try:
            return Decimal(text)
        except (
            InvalidOperation,
            ValueError,
        ):
            return None

    @staticmethod
    def _to_int(
        value: object,
    ) -> int | None:
        if value is None:
            return None

        if isinstance(
            value,
            bool,
        ):
            return int(value)

        if isinstance(
            value,
            int,
        ):
            return value

        if isinstance(
            value,
            float,
        ):
            return int(value)

        text = str(value).strip()

        if not text:
            return None

        try:
            return int(Decimal(text))
        except (
            InvalidOperation,
            ValueError,
        ):
            return None

    @staticmethod
    def _decimal_equal(
        left: Decimal | None,
        right: Decimal | None,
    ) -> bool:
        if left is None and right is None:
            return True

        if left is None or right is None:
            return False

        return left == right

    # ---------------------------------------------------------
    # Result helpers
    # ---------------------------------------------------------

    @staticmethod
    def _change_to_dict(
        change: ProductChange,
    ) -> dict[str, object]:
        return {
            "supplier_sku": (change.supplier_sku),
            "row_number": (change.row_number),
            "is_new": change.is_new,
            "cost_changed": (change.cost_changed),
            "availability_changed": (change.availability_changed),
            "name_changed": (change.name_changed),
            "url_changed": (change.url_changed),
            "existing_cost": (
                str(change.existing_cost) if change.existing_cost is not None else None
            ),
            "api_cost": (str(change.api_cost) if change.api_cost is not None else None),
            "existing_availability": (change.existing_availability),
            "api_availability": (change.api_availability),
            "existing_name": (change.existing_name),
            "api_name": (change.api_name),
            "existing_url": (change.existing_url),
            "api_url": (change.api_url),
        }

    def _create_backup_path(
        self,
        *,
        label: str,
    ) -> Path:
        """
        Generate a unique timestamped workbook backup path.
        """

        timestamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")

        return self.workbook_path.with_name(
            f"{self.workbook_path.stem}."
            f"{label}."
            f"{timestamp}"
            f"{self.workbook_path.suffix}"
        )
