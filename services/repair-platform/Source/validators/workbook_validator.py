from pathlib import Path

from openpyxl import load_workbook

from config.database import TABLES


class WorkbookValidator:

    def __init__(self, workbook_path):

        self.workbook_path = Path(workbook_path)

    def validate(self):

        print()
        print("=" * 70)
        print("WORKBOOK VALIDATION")
        print("=" * 70)

        # Check that workbook exists
        if not self.workbook_path.exists():

            print(f"✗ Workbook not found:\n{self.workbook_path}")
            return False

        print("✓ Workbook Found")

        # Open workbook
        try:

            workbook = load_workbook(
                self.workbook_path,
                data_only=True
            )

        except Exception as error:

            print("✗ Unable to open workbook")
            print(error)
            return False

        print("✓ Workbook Opened")

        # Collect every table in the workbook
        workbook_tables = set()

        for worksheet in workbook.worksheets:

            for table in worksheet.tables.values():

                workbook_tables.add(table.name)

        print()
        print("Checking Required Tables")
        print("-" * 70)

        missing = []

        for table_name in TABLES.values():

            if table_name in workbook_tables:

                print(f"✓ {table_name}")

            else:

                print(f"✗ {table_name}")
                missing.append(table_name)

        print("-" * 70)
        print(f"Total Tables Found : {len(workbook_tables)}")
        print(f"Required Tables    : {len(TABLES)}")

        if missing:

            print()
            print("Missing Tables:")

            for table in missing:

                print(f" • {table}")

            print()
            print("Workbook Validation FAILED")

            return False

        print()
        print("Workbook Validation PASSED")

        return True