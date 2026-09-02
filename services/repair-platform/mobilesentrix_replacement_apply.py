from __future__ import annotations

import json
import shutil
from dataclasses import asdict, dataclass
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

import integrations.mobilesentrix.workbook_sync as workbook_sync
from integrations.mobilesentrix import (
    MobileSentrixWorkbookSyncService,
)
from integrations.mobilesentrix.client import (
    MobileSentrixClient,
)
from integrations.mobilesentrix.models import (
    MobileSentrixDetailedProduct,
)

type ExcelCellValue = str | int | float | bool | datetime | None
# ============================================================
# SAFETY CONFIGURATION
# ============================================================

# FIRST RUN MUST REMAIN FALSE.
#
# False:
#     Validate everything and generate a report.
#     No workbook or checkpoint changes.
#
# True:
#     Apply only the explicitly approved mappings below,
#     after all validation passes.
#
APPLY_CHANGES = False


# ============================================================
# Approved replacements
# ============================================================

APPROVED_REPLACEMENTS = [
    {
        "row": 64,
        "old_sku": "107082133697",
        "new_sku": "107182127725",
        "product_id": "249690",
        "score": 0.9447,
        "note": (
            "Same Aixun iHeater Pro 4th Gen North American "
            "110V product; compatibility expanded through "
            "iPhone 17 Pro Max."
        ),
    },
    {
        "row": 7436,
        "old_sku": "107282225908",
        "new_sku": "107082069206",
        "product_id": "90349",
        "score": 1.0,
        "note": (
            "Same iPhone 11 Pro Max aftermarket Incell "
            "display family; current listing identified as "
            "AQ7 / Incell."
        ),
    },
    {
        "row": 37280,
        "old_sku": "107183202027",
        "new_sku": "107183202032",
        "product_id": "254905",
        "score": 1.0,
        "note": (
            "Identical Xiaomi Redmi Note 10 Service Pack "
            "display listing under current Mobile Sentrix SKU."
        ),
    },
]


# ============================================================
# Paths
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

REPORT_DIR = BASE_DIR / "Data" / "reports"

RUN_TIMESTAMP = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")

REPORT_PATH = REPORT_DIR / f"mobilesentrix_replacement_apply_{RUN_TIMESTAMP}.json"


# ============================================================
# Result model
# ============================================================


@dataclass(slots=True)
class ReplacementResult:
    row: int

    expected_old_sku: str
    proposed_new_sku: str
    product_id: str
    approved_score: float

    workbook_current_sku: str | None

    live_sku: str | None
    live_name: str | None
    live_price: str | None

    live_in_stock: bool | None
    live_stock_quantity: int | None
    live_saleable: bool | None
    live_end_of_life: bool | None
    live_product_url: str | None

    checkpoint_status_before: str | None

    workbook_validation_passed: bool
    api_validation_passed: bool
    checkpoint_validation_passed: bool

    eligible_to_apply: bool
    applied: bool

    validation_messages: list[str]


# ============================================================
# General helpers
# ============================================================


def normalize_sku(
    value: object,
) -> str:
    return str(value or "").strip()


def safe_string(
    value: object,
) -> str | None:
    if value is None:
        return None

    text = str(value).strip()

    return text or None


def decimal_string(
    value: Decimal | None,
) -> str | None:
    if value is None:
        return None

    return format(
        value,
        "f",
    )


def utc_now() -> str:
    return datetime.now(UTC).isoformat()


# ============================================================
# Workbook helpers
# ============================================================
def set_cell_value(
    *,
    worksheet: Any,
    row: int,
    column: int,
    value: ExcelCellValue,
) -> None:
    """
    Safely write a supported value to a normal workbook cell.

    openpyxl may represent cells inside merged ranges as
    MergedCell objects. Those cells are placeholders and
    cannot be written to.

    This helper also restricts values to the scalar types
    written by the Mobile Sentrix replacement process.
    """

    cell = worksheet.cell(
        row=row,
        column=column,
    )

    if not isinstance(
        cell,
        Cell,
    ):
        raise RuntimeError(
            "Refusing to write to a non-writable or merged cell: "
            f"row={row}, column={column}"
        )

    cell.value = value


def build_header_map(
    worksheet: Any,
) -> dict[str, int]:
    headers: dict[str, int] = {}

    for column in range(
        1,
        worksheet.max_column + 1,
    ):
        value = worksheet.cell(
            1,
            column,
        ).value

        if value is None:
            continue

        header = str(value).strip()

        if header:
            headers[header] = column

    return headers


def require_headers(
    headers: dict[str, int],
) -> None:
    required = {
        "Supplier SKU",
        "Product Name",
        "Unit Cost",
        "Availability",
        "Product URL",
        "Observed Date",
        "Source File",
    }

    missing = required - set(headers)

    if missing:
        raise RuntimeError(
            "Required workbook columns are missing: " + ", ".join(sorted(missing))
        )


# ============================================================
# Checkpoint helpers
# ============================================================


def find_checkpoint_item(
    *,
    retry_items: list[object],
    row_number: int,
    sku: str,
) -> dict[str, object] | None:
    for raw_item in retry_items:
        if not isinstance(
            raw_item,
            dict,
        ):
            continue

        try:
            item_row = int(str(raw_item.get("row")))

        except (
            TypeError,
            ValueError,
        ):
            continue

        item_sku = normalize_sku(raw_item.get("sku"))

        if item_row == row_number and item_sku == sku:
            return raw_item

    return None


# ============================================================
# Live product validation
# ============================================================


def load_live_product(
    *,
    client: MobileSentrixClient,
    product_id: str,
) -> MobileSentrixDetailedProduct:
    raw = client.get_product(product_id=product_id)

    return MobileSentrixDetailedProduct.from_api_item(raw)


# ============================================================
# Main
# ============================================================


def main() -> int:
    REPORT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    print()
    print("=" * 72)
    print("Mobile Sentrix Approved Replacement Apply")
    print("=" * 72)

    print(
        "Mode:",
        ("APPLY" if APPLY_CHANGES else "DRY RUN"),
    )

    print(
        "Approved replacements:",
        len(APPROVED_REPLACEMENTS),
    )

    sync = MobileSentrixWorkbookSyncService()

    client = MobileSentrixClient()

    workbook_path = sync.workbook_path

    checkpoint = sync.get_sync_checkpoint()

    retry_items = checkpoint.get(
        "retry_items",
        [],
    )

    if not isinstance(
        retry_items,
        list,
    ):
        raise RuntimeError("Checkpoint retry_items is not a list.")

    keep_vba = workbook_path.suffix.lower() == ".xlsm"

    workbook = load_workbook(
        workbook_path,
        read_only=False,
        data_only=False,
        keep_vba=keep_vba,
    )

    results: list[ReplacementResult] = []

    backup_path: Path | None = None

    workbook_changed = False
    checkpoint_changed = False

    try:
        if workbook_sync.SHEET_NAME not in workbook.sheetnames:
            raise RuntimeError("Worksheet not found: " f"{workbook_sync.SHEET_NAME}")

        worksheet = workbook[workbook_sync.SHEET_NAME]

        headers = build_header_map(worksheet)

        require_headers(headers)

        for index, approved in enumerate(
            APPROVED_REPLACEMENTS,
            start=1,
        ):
            row_number = int(approved["row"])

            old_sku = normalize_sku(approved["old_sku"])

            new_sku = normalize_sku(approved["new_sku"])

            product_id = str(approved["product_id"]).strip()

            approved_score = float(approved["score"])

            messages: list[str] = []

            print()
            print("-" * 72)

            print(f"Replacement " f"{index}/" f"{len(APPROVED_REPLACEMENTS)}")

            print(
                "Workbook row:",
                row_number,
            )

            print(
                "Old SKU:",
                old_sku,
            )

            print(
                "New SKU:",
                new_sku,
            )

            # ------------------------------------------------
            # Workbook validation
            # ------------------------------------------------

            workbook_current_sku = normalize_sku(
                worksheet.cell(
                    row_number,
                    headers["Supplier SKU"],
                ).value
            )

            workbook_validation_passed = workbook_current_sku == old_sku

            if workbook_validation_passed:
                messages.append("Workbook row contains expected old SKU.")

            else:
                messages.append(
                    "WORKBOOK SKU MISMATCH: "
                    f"expected {old_sku}, "
                    f"found {workbook_current_sku}."
                )

            # ------------------------------------------------
            # Checkpoint validation
            # ------------------------------------------------

            checkpoint_item = find_checkpoint_item(
                retry_items=retry_items,
                row_number=row_number,
                sku=old_sku,
            )

            checkpoint_status_before: str | None = None

            checkpoint_validation_passed = False

            if checkpoint_item is None:
                messages.append("Matching checkpoint record was not found.")

            else:
                checkpoint_status_before = (
                    str(
                        checkpoint_item.get(
                            "status",
                            "",
                        )
                    )
                    .strip()
                    .lower()
                )

                checkpoint_validation_passed = checkpoint_status_before == "unmatched"

                if checkpoint_validation_passed:
                    messages.append("Checkpoint record is currently unmatched.")

                else:
                    messages.append(
                        "CHECKPOINT STATUS MISMATCH: " f"{checkpoint_status_before}"
                    )

            # ------------------------------------------------
            # Live Mobile Sentrix validation
            # ------------------------------------------------

            live_product = load_live_product(
                client=client,
                product_id=product_id,
            )

            live_sku = normalize_sku(live_product.sku)

            api_validation_passed = live_sku == new_sku

            if api_validation_passed:
                messages.append(
                    "Live Mobile Sentrix product SKU " "matches approved new SKU."
                )

            else:
                messages.append(
                    "LIVE API SKU MISMATCH: "
                    f"expected {new_sku}, "
                    f"received {live_sku}."
                )

            if live_product.end_of_life is True:
                messages.append("WARNING: Candidate is marked end-of-life.")

            if live_product.is_saleable is False:
                messages.append("WARNING: Candidate is not currently saleable.")

            # ------------------------------------------------
            # Final eligibility
            # ------------------------------------------------

            eligible_to_apply = (
                workbook_validation_passed
                and checkpoint_validation_passed
                and api_validation_passed
            )

            applied = False

            # ------------------------------------------------
            # Apply
            # ------------------------------------------------

            if APPLY_CHANGES and eligible_to_apply:
                observed_at = utc_now()

                availability = (
                    "In Stock"
                    if live_product.is_in_stock is True
                    else (
                        "Out of Stock"
                        if live_product.is_in_stock is False
                        else "Unknown"
                    )
                )

                set_cell_value(
                    worksheet=worksheet,
                    row=row_number,
                    column=headers["Supplier SKU"],
                    value=new_sku,
                )

                if live_product.name:
                    set_cell_value(
                        worksheet=worksheet,
                        row=row_number,
                        column=headers["Product Name"],
                        value=live_product.name,
                    )

                if live_product.customer_price is not None:
                    set_cell_value(
                        worksheet=worksheet,
                        row=row_number,
                        column=headers["Unit Cost"],
                        value=float(live_product.customer_price),
                    )

                set_cell_value(
                    worksheet=worksheet,
                    row=row_number,
                    column=headers["Availability"],
                    value=availability,
                )

                if live_product.product_url:
                    set_cell_value(
                        worksheet=worksheet,
                        row=row_number,
                        column=headers["Product URL"],
                        value=(live_product.product_url),
                    )

                set_cell_value(
                    worksheet=worksheet,
                    row=row_number,
                    column=headers["Observed Date"],
                    value=observed_at,
                )

                set_cell_value(
                    worksheet=worksheet,
                    row=row_number,
                    column=headers["Source File"],
                    value=("Mobile Sentrix API " "Replacement Resolution"),
                )

                if checkpoint_item is None:
                    raise RuntimeError(
                        "Checkpoint item disappeared " "during replacement processing."
                    )

                checkpoint_item["status"] = "resolved_replacement"

                checkpoint_item["resolution_type"] = "replacement_sku"

                checkpoint_item["resolved_at"] = observed_at

                checkpoint_item["old_sku"] = old_sku

                checkpoint_item["replacement_sku"] = new_sku

                checkpoint_item["replacement_product_id"] = product_id

                checkpoint_item["resolution_score"] = approved_score

                checkpoint_item["resolution_note"] = approved.get("note")

                checkpoint_item["resolved_product_name"] = live_product.name

                checkpoint_item["resolved_price"] = decimal_string(
                    live_product.customer_price
                )

                checkpoint_item["resolved_in_stock"] = live_product.is_in_stock

                checkpoint_item["resolved_stock_quantity"] = live_product.in_stock_qty

                workbook_changed = True
                checkpoint_changed = True
                applied = True

                messages.append("Replacement applied.")

            elif APPLY_CHANGES and not eligible_to_apply:
                messages.append("NOT APPLIED because validation failed.")

            else:
                if eligible_to_apply:
                    messages.append("DRY RUN: eligible to apply.")

                else:
                    messages.append("DRY RUN: not eligible to apply.")

            result = ReplacementResult(
                row=row_number,
                expected_old_sku=(old_sku),
                proposed_new_sku=(new_sku),
                product_id=(product_id),
                approved_score=(approved_score),
                workbook_current_sku=(workbook_current_sku),
                live_sku=(safe_string(live_product.sku)),
                live_name=(live_product.name),
                live_price=(decimal_string(live_product.customer_price)),
                live_in_stock=(live_product.is_in_stock),
                live_stock_quantity=(live_product.in_stock_qty),
                live_saleable=(live_product.is_saleable),
                live_end_of_life=(live_product.end_of_life),
                live_product_url=(live_product.product_url),
                checkpoint_status_before=(checkpoint_status_before),
                workbook_validation_passed=(workbook_validation_passed),
                api_validation_passed=(api_validation_passed),
                checkpoint_validation_passed=(checkpoint_validation_passed),
                eligible_to_apply=(eligible_to_apply),
                applied=(applied),
                validation_messages=(messages),
            )

            results.append(result)

            print(
                "Workbook validation:",
                workbook_validation_passed,
            )

            print(
                "Checkpoint validation:",
                checkpoint_validation_passed,
            )

            print(
                "API validation:",
                api_validation_passed,
            )

            print(
                "Live product:",
                live_product.name,
            )

            print(
                "Live price:",
                live_product.customer_price,
            )

            print(
                "Live stock:",
                live_product.in_stock_qty,
            )

            print(
                "Eligible to apply:",
                eligible_to_apply,
            )

            print(
                "Applied:",
                applied,
            )

            for message in messages:
                print(
                    "  -",
                    message,
                )

        # ====================================================
        # Save changes
        # ====================================================

        if APPLY_CHANGES:
            eligible_count = sum(1 for item in results if item.eligible_to_apply)

            applied_count = sum(1 for item in results if item.applied)

            if eligible_count != len(APPROVED_REPLACEMENTS):
                raise RuntimeError(
                    "Not all approved replacements "
                    "passed validation. "
                    "No files will be saved."
                )

            if applied_count != len(APPROVED_REPLACEMENTS):
                raise RuntimeError(
                    "Not all approved replacements "
                    "were applied. "
                    "No files will be saved."
                )

            # ------------------------------------------------
            # Backup workbook before save
            # ------------------------------------------------

            backup_path = workbook_path.with_name(
                workbook_path.stem
                + ".before-approved-replacements-"
                + RUN_TIMESTAMP
                + workbook_path.suffix
            )

            shutil.copy2(
                workbook_path,
                backup_path,
            )

            print()
            print(
                "Workbook backup:",
                backup_path,
            )

            # ------------------------------------------------
            # Save workbook
            # ------------------------------------------------

            if workbook_changed:
                workbook.save(workbook_path)

            # ------------------------------------------------
            # Save checkpoint
            # ------------------------------------------------

            if checkpoint_changed:
                checkpoint["last_resolution_run_at"] = utc_now()

                pending_count = sum(
                    1
                    for item in retry_items
                    if (
                        isinstance(
                            item,
                            dict,
                        )
                        and str(
                            item.get(
                                "status",
                                "",
                            )
                        )
                        .strip()
                        .lower()
                        == "pending"
                    )
                )

                checkpoint["cycle_complete"] = (
                    checkpoint.get("next_start_row") is None and pending_count == 0
                )

                sync._save_checkpoint(checkpoint)

    finally:
        workbook.close()

    # ========================================================
    # Report
    # ========================================================

    eligible_count = sum(1 for item in results if item.eligible_to_apply)

    applied_count = sum(1 for item in results if item.applied)

    report = {
        "generated_at": (utc_now()),
        "mode": ("apply" if APPLY_CHANGES else "dry_run"),
        "workbook": str(workbook_path),
        "approved_replacement_count": (len(APPROVED_REPLACEMENTS)),
        "eligible_count": (eligible_count),
        "applied_count": (applied_count),
        "backup_path": (str(backup_path) if backup_path is not None else None),
        "results": [asdict(item) for item in results],
    }

    REPORT_PATH.write_text(
        json.dumps(
            report,
            indent=2,
            default=str,
        ),
        encoding="utf-8",
    )

    print()
    print("=" * 72)

    print(
        "Replacement validation complete"
        if not APPLY_CHANGES
        else "Replacement apply complete"
    )

    print("=" * 72)

    print(
        "Mode:",
        ("DRY RUN" if not APPLY_CHANGES else "APPLY"),
    )

    print(
        "Approved:",
        len(APPROVED_REPLACEMENTS),
    )

    print(
        "Eligible:",
        eligible_count,
    )

    print(
        "Applied:",
        applied_count,
    )

    print(
        "Report:",
        REPORT_PATH,
    )

    if not APPLY_CHANGES:
        print()
        print("No workbook or checkpoint changes were made.")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
