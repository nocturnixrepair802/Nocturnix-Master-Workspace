from __future__ import annotations

import shutil
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from integrations.mobilesentrix import MobileSentrixWorkbookSyncService


def retry_count(checkpoint: dict[str, object]) -> int:
    value = checkpoint.get("retry_items", [])
    if not isinstance(value, list):
        return 0
    return len(value)


def main() -> None:
    sync = MobileSentrixWorkbookSyncService()
    checkpoint_path = Path(sync._checkpoint_path())
    backup = checkpoint_path.with_name(checkpoint_path.stem + ".before_retry_test.json")

    print("\nMobile Sentrix Controlled Retry Recovery Test")
    print("---------------------------------------------")
    shutil.copy2(checkpoint_path, backup)
    print("Checkpoint backup:", backup)

    checkpoint = sync.get_sync_checkpoint()
    row_number = checkpoint.get("next_start_row")
    if not isinstance(row_number, int):
        raise RuntimeError(f"Expected integer next_start_row, got {row_number!r}")

    keep_vba = sync.workbook_path.suffix.lower() == ".xlsm"
    wb = load_workbook(
        sync.workbook_path,
        read_only=True,
        data_only=False,
        keep_vba=keep_vba,
    )
    try:
        ws = wb["40 - Parts"]
        headers = {
            str(cell.value).strip(): cell.column
            for cell in ws[1]
            if cell.value is not None
        }
        sku_value = ws.cell(
            row=row_number,
            column=headers["Supplier SKU"],
        ).value
        supplier_value = ws.cell(
            row=row_number,
            column=headers["Supplier"],
        ).value
    finally:
        wb.close()

    sku = str(sku_value).strip() if sku_value is not None else ""
    supplier = str(supplier_value).strip() if supplier_value is not None else ""

    print("Test row:", row_number)
    print("Supplier:", supplier)
    print("SKU:", sku)

    if supplier != "MobileSentrix":
        raise RuntimeError(f"Row {row_number} is not a MobileSentrix row.")
    if not sku:
        raise RuntimeError(f"Row {row_number} has no Supplier SKU.")

    original_next = checkpoint.get("next_start_row")
    original_processed = checkpoint.get("products_processed_this_cycle")

    retry_items_value = checkpoint.get("retry_items", [])
    retry_items: list[dict[str, object]] = []
    if isinstance(retry_items_value, list):
        for raw_item in retry_items_value:
            if isinstance(raw_item, dict):
                retry_items.append({str(key): value for key, value in raw_item.items()})

    now = datetime.now(UTC).isoformat()
    retry_items.append(
        {
            "row": row_number,
            "sku": sku,
            "last_error": "Controlled retry recovery test",
            "attempts": 0,
            "first_failed_at": now,
            "last_failed_at": now,
            "status": "pending",
        }
    )
    checkpoint["retry_items"] = retry_items
    checkpoint["cycle_complete"] = False
    sync._save_checkpoint(checkpoint)

    before = sync.get_sync_checkpoint()
    print("\nRetry item inserted.")
    print(
        "Retry count before processing:",
        retry_count(before),
    )

    result = sync.retry_failed_items(
        max_items=1,
        request_delay_seconds=0.25,
    )

    print("\nRetry Processor Result")
    print("----------------------")
    keys = [
        "retry_items_processed",
        "recovered",
        "still_pending",
        "exhausted",
        "api_errors",
        "rows_updated",
        "unchanged_rows",
        "retry_count",
        "supplier_data_changed",
        "observation_metadata_refreshed",
        "workbook_saved",
    ]
    for key in keys:
        print(f"{key}: {result.get(key)}")

    final = sync.get_sync_checkpoint()
    print("\nCheckpoint Verification")
    print("-----------------------")
    print("Next start row:", final.get("next_start_row"))
    print(
        "Products processed:",
        final.get("products_processed_this_cycle"),
    )
    print(
        "Retry count:",
        retry_count(final),
    )
    print("Cycle complete:", final.get("cycle_complete"))

    if final.get("next_start_row") != original_next:
        raise RuntimeError("FAIL: retry processing changed next_start_row.")
    if final.get("products_processed_this_cycle") != original_processed:
        raise RuntimeError("FAIL: retry processing changed processed count.")

    print("\nPrimary checkpoint position preserved: PASS")
    if result.get("recovered") == 1:
        print("Controlled retry recovery: PASS")
    else:
        print("Controlled retry recovery: REVIEW REQUIRED")

    print("\nCheckpoint backup retained at:")
    print(backup)


if __name__ == "__main__":
    main()
